import os
import io
import re
import copy
from datetime import datetime
import docx
from docx.shared import Pt
from config import OFFICIAL_UNITS, logger

MONTHS_ID = {
    'JANUARI': '01', 'FEBRUARI': '02', 'MARET': '03', 'APRIL': '04',
    'MEI': '05', 'JUNI': '06', 'JULI': '07', 'AGUSTUS': '08',
    'SEPTEMBER': '09', 'OKTOBER': '10', 'NOVEMBER': '11', 'DESEMBER': '12'
}

DAYS_ID = {
    0: 'SENIN', 1: 'SELASA', 2: 'RABU', 3: 'KAMIS', 4: 'JUMAT', 5: 'SABTU', 6: 'MINGGU'
}

def clean_text(text: str) -> str:
    if not text:
        return ''
    t = text.replace('\xa0', ' ').replace('–', '-').replace('—', '-')
    t = re.sub(r'\s+', ' ', t)
    return t.strip()

def extract_date_from_text(text: str) -> str | None:
    match = re.search(r'(\d{1,2})\s+(JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+(\d{4})', text.upper())
    if match:
        day = match.group(1).zfill(2)
        month = MONTHS_ID[match.group(2)]
        year = match.group(3)
        return f"{year}-{month}-{day}"
    match_iso = re.search(r'(\d{4})-(\d{2})-(\d{2})', text)
    if match_iso:
        return match_iso.group(0)
    return None

def format_indonesian_date(date_str: str) -> str:
    try:
        dt = datetime.strptime(date_str, "%Y-%m-%d")
        day_name = DAYS_ID[dt.weekday()]
        month_keys = list(MONTHS_ID.keys())
        month_name = month_keys[dt.month - 1]
        return f"{day_name}, {dt.day:02d} {month_name} {dt.year}"
    except Exception:
        return date_str

def get_row_cells(row):
    seen = set()
    cells = []
    for c in row.cells:
        if c._tc not in seen:
            seen.add(c._tc)
            cells.append(clean_text(c.text))
    return cells

def set_cell_text(cell, text, bold=False, font_size=12.0, font_name='Arial'):
    p = cell.paragraphs[0]
    p.text = ''
    run = p.add_run(str(text if text is not None else ''))
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold

def detect_subsatgas(filename: str, doc: docx.Document) -> str:
    upper_fname = filename.upper()
    if "MINOPS" in upper_fname:
        return "MINOPS"
    if "LIDIK" in upper_fname:
        return "LIDIK"
    if "INVEST" in upper_fname:
        return "INVESTIGASI"
    if "ANALIS" in upper_fname:
        return "ANALIS"

    for p in doc.paragraphs[:10]:
        t = p.text.upper()
        if "SUBSATGAS ANALIS" in t:
            return "ANALIS"
        if "SUBSATGAS LIDIK" in t:
            return "LIDIK"
        if "SUBSATGAS INVEST" in t:
            return "INVESTIGASI"
        if "MINOPS" in t:
            return "MINOPS"

    return "MINOPS"

def normalize_unit_name(unit_raw: str, subsatgas: str) -> str:
    clean = clean_text(unit_raw).upper()
    clean = re.sub(r'^(?:[IVXLCDM]+|\d+)[.)]?\s*', '', clean)
    clean = re.sub(r'^(?:UNIT|WILAYAH)\s+', '', clean)
    clean = clean.replace('–', '-').replace('—', '-').strip()

    if 'PEGUBIN' in clean or 'PEGUNUNG BINTANG' in clean or 'PEGUNUNGAN BINTANG' in clean:
        clean = 'PEGUNUNGAN BINTANG'
    elif 'WAMENA' in clean or 'KRING JAYAWIJAYA' in clean:
        clean = 'LANNY JAYA - KRING JAYAWIJAYA'
    elif 'SINAK' in clean:
        clean = 'PUNCAK KRING SINAK'
    elif 'ILAGA' in clean or 'PUNCAK / ILAGA' in clean:
        clean = 'PUNCAK / ILAGA'

    prefix = "WILAYAH " if subsatgas == "ANALIS" else "UNIT "
    formatted = prefix + clean

    official_list = OFFICIAL_UNITS.get(subsatgas, [])
    for off in official_list:
        off_clean = off.replace('–', '-').replace('—', '-').upper()
        if formatted == off_clean:
            return off
        if clean in off_clean:
            return off

    return formatted

def parse_docx_bytes(file_bytes: bytes, filename: str) -> dict:
    doc = docx.Document(io.BytesIO(file_bytes))
    subsatgas = detect_subsatgas(filename, doc)
    
    detected_date = extract_date_from_text(filename)
    if not detected_date:
        for p in doc.paragraphs[:15]:
            d = extract_date_from_text(p.text)
            if d:
                detected_date = d
                break

    items = []
    incidents = []

    if not doc.tables:
        return {
            "filename": filename,
            "subsatgas": subsatgas,
            "detected_date": detected_date,
            "items": items,
            "incidents": incidents
        }

    tbl = doc.tables[0]
    
    default_unit = None
    if subsatgas == "MINOPS":
        upper_fname = filename.upper()
        if "JAYAPURA" in upper_fname:
            default_unit = "UNIT JAYAPURA"
        elif "TIMIKA" in upper_fname:
            default_unit = "UNIT TIMIKA"
        elif "JAYAWIJAYA" in upper_fname or "WAMENA" in upper_fname or "LANNY" in upper_fname:
            default_unit = "UNIT LANNY JAYA – KRING JAYAWIJAYA"

    current_unit = default_unit
    last_context = {"time": "", "location": "", "personnel": ""}

    for row in tbl.rows:
        cells = get_row_cells(row)
        if not any(cells):
            continue

        non_empty = [c for c in cells if c]
        if len(non_empty) == 1:
            val = non_empty[0].upper()
            if any(k in val for k in ['UNIT', 'WILAYAH', 'KRING', 'JAYAPURA', 'TIMIKA', 'LANNY', 'YAHUKIMO', 'NDUGA', 'INTAN', 'NABIRE', 'PUNCAK', 'MIMIKA', 'PEGUNUNG']):
                current_unit = normalize_unit_name(non_empty[0], subsatgas)
                last_context = {"time": "", "location": "", "personnel": ""}
                continue

        row_str = ' '.join(cells).upper()
        if 'NO' in cells[0].upper() and ('WAKTU' in row_str or 'SATGAS' in row_str or 'JAM' in row_str or 'PUKUL' in row_str):
            continue
        if row_str.replace(' ', '') in ['1234567', '12345678']:
            continue
        if 'KEPOLISIAN NEGARA' in row_str or 'SATUAN TUGAS' in row_str or 'KASUBSATGAS' in row_str:
            continue

        if len(cells) < 6:
            continue

        item = {
            "subsatgas": subsatgas,
            "unit_name": current_unit or "UNIT JAYAPURA",
            "satgas": "Gakkum",
            "activity_time": "",
            "personnel_count": "",
            "location": "",
            "activity_desc": "",
            "results": "",
            "notes": ""
        }

        if subsatgas == "MINOPS":
            item["satgas"] = cells[1] if len(cells) > 1 and cells[1] else "Gakkum"
            item["activity_time"] = cells[2] if len(cells) > 2 else ""
            item["personnel_count"] = cells[3] if len(cells) > 3 else ""
            item["location"] = cells[4] if len(cells) > 4 else ""
            item["activity_desc"] = cells[5] if len(cells) > 5 else ""
            item["results"] = cells[6] if len(cells) > 6 else ""
            item["notes"] = cells[7] if len(cells) > 7 else ""

        elif subsatgas in ["LIDIK", "INVESTIGASI"]:
            item["activity_time"] = cells[1] if len(cells) > 1 else ""
            item["personnel_count"] = cells[2] if len(cells) > 2 else ""
            item["location"] = cells[3] if len(cells) > 3 else ""
            item["activity_desc"] = cells[4] if len(cells) > 4 else ""
            item["results"] = cells[5] if len(cells) > 5 else ""
            item["notes"] = cells[6] if len(cells) > 6 else ""

        elif subsatgas == "ANALIS":
            item["activity_time"] = cells[1] if len(cells) > 1 else ""
            item["location"] = cells[2] if len(cells) > 2 else ""
            item["activity_desc"] = cells[3] if len(cells) > 3 else ""
            item["results"] = cells[4] if len(cells) > 4 else ""
            item["personnel_count"] = cells[5] if len(cells) > 5 else ""
            item["notes"] = cells[6] if len(cells) > 6 else ""

            if item["activity_time"]:
                last_context["time"] = item["activity_time"]
            else:
                item["activity_time"] = last_context["time"]

            if item["location"]:
                last_context["location"] = item["location"]
            else:
                item["location"] = last_context["location"]

            if item["personnel_count"]:
                last_context["personnel"] = item["personnel_count"]
            else:
                item["personnel_count"] = last_context["personnel"]

        if item["activity_desc"] or item["results"]:
            items.append(item)

    if len(doc.tables) > 1:
        inc_tbl = doc.tables[1]
        for r in inc_tbl.rows[1:]:
            c = get_row_cells(r)
            if len(c) >= 6 and (c[3] or c[2]):
                if c[3].strip() not in ['-', '', '=']:
                    incidents.append({
                        "incident_time": c[1] if len(c) > 1 else "",
                        "location": c[2] if len(c) > 2 else "",
                        "incident_desc": c[3] if len(c) > 3 else "",
                        "casualties_losses": c[4] if len(c) > 4 else "",
                        "actions_taken": c[5] if len(c) > 5 else "",
                        "notes": c[6] if len(c) > 6 else ""
                    })

    return {
        "filename": filename,
        "subsatgas": subsatgas,
        "detected_date": detected_date,
        "items": items,
        "incidents": incidents
    }

def merge_items_to_master(template_path: str, items: list[dict], incidents: list[dict], target_date: str, kasatgas_config: dict = None) -> bytes:
    doc = docx.Document(template_path)
    formatted_date_upper = format_indonesian_date(target_date).upper()

    # 1. Update Title Date in Paragraphs
    for p in doc.paragraphs[:20]:
        if "PERIODE HARI/TANGGAL" in p.text.upper():
            p.text = f"PERIODE HARI/TANGGAL : {formatted_date_upper}"
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9.5)
                r.bold = True

    # 2. Update Signature Block and prevent splitting across pages
    city = (kasatgas_config or {}).get("signature_city", "Jayapura")
    kasatgas_name = (kasatgas_config or {}).get("kasatgas_name", "I G.G. ERA ADHINATA, S.I.K.")
    kasatgas_rank = (kasatgas_config or {}).get("kasatgas_rank_nrp", "KOMISARIS BESAR POLISI NRP 80020993")

    parts = formatted_date_upper.split(',')
    date_only = parts[-1].strip() if len(parts) > 1 else formatted_date_upper

    # Find start of signature block
    sig_start_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if re.search(r'JAYAPURA,\s+\d{1,2}\s+\w+\s+\d{4}', p.text.upper()) or (city.upper() in p.text.upper() and re.search(r'\d{4}', p.text)):
            sig_start_idx = idx
            break

    # Clean up excess empty paragraphs before signature block (keep at most 1 empty line)
    if sig_start_idx is not None:
        check_idx = sig_start_idx - 1
        empty_count = 0
        while check_idx >= 0 and not doc.paragraphs[check_idx].text.strip():
            empty_count += 1
            if empty_count > 1 and not doc.paragraphs[check_idx]._p.xpath('.//w:drawing') and not doc.paragraphs[check_idx]._p.xpath('.//w:pict'):
                doc.paragraphs[check_idx]._p.getparent().remove(doc.paragraphs[check_idx]._p)
            check_idx -= 1

    # Scan and format signature block paragraphs
    sig_paragraphs = []
    is_in_sig = False
    for p in doc.paragraphs:
        upper_text = p.text.upper()
        if re.search(r'JAYAPURA,\s+\d{1,2}\s+\w+\s+\d{4}', upper_text) or (city.upper() in upper_text and re.search(r'\d{4}', upper_text)):
            is_in_sig = True
            p.text = f"{city}, {date_only.title()}"
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(12.0)
            sig_paragraphs.append(p)
            continue

        if is_in_sig:
            sig_paragraphs.append(p)
            if "KOMISARIS BESAR POLISI" in upper_text or (kasatgas_config and "NRP" in upper_text):
                p.text = kasatgas_rank
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(12.0)
                break
            elif "ERA ADHINATA" in upper_text or (kasatgas_config and kasatgas_name.upper() in upper_text):
                p.text = kasatgas_name
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(12.0)
                    r.bold = True
                    r.underline = True
            elif "KEPALA SATUAN TUGAS" in upper_text or "OPERASI DAMAI CARTENZ" in upper_text:
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(12.0)

    # Clean up redundant empty lines between title and name (keep at most 1 spacing)
    blank_paragraphs = [p for p in sig_paragraphs if not p.text.strip() and not p._p.xpath('.//w:drawing') and not p._p.xpath('.//w:pict')]
    if len(blank_paragraphs) > 1:
        for bp in blank_paragraphs[1:]:
            bp._p.getparent().remove(bp._p)
            if bp in sig_paragraphs:
                sig_paragraphs.remove(bp)

    # Bind all signature paragraphs together with keep_with_next so they NEVER split across pages!
    for i, p in enumerate(sig_paragraphs):
        p.paragraph_format.keep_lines = True
        if i < len(sig_paragraphs) - 1:
            p.paragraph_format.keep_with_next = True

    # 3. Group items by (subsatgas, unit_name)
    items_by_unit = {}
    for it in items:
        sub = it.get('subsatgas', 'MINOPS').upper()
        unit = it.get('unit_name', '').upper().replace('–', '-').replace('—', '-')
        key = (sub, unit)
        items_by_unit.setdefault(key, []).append(it)

    # 4. Map template Table 0 sections
    t0 = doc.tables[0]
    sections = []
    current_sub = 'MINOPS'

    for idx, r in enumerate(t0.rows):
        seen = set()
        unique = [clean_text(c.text).upper() for c in r.cells if not (c._tc in seen or seen.add(c._tc))]
        non_empty = [c for c in unique if c]
        if len(non_empty) == 1:
            txt = non_empty[0]
            if txt in ['MINOPS', 'SUBSATGAS LIDIK', 'SUBSATGAS INVESTIGASI', 'SUBSATGAS ANALIS']:
                current_sub = txt.replace('SUBSATGAS ', '')
                continue
            elif any(k in txt for k in ['UNIT', 'WILAYAH']):
                sections.append({
                    'subsatgas': current_sub,
                    'unit_name': txt.replace('–', '-').replace('—', '-'),
                    'heading_idx': idx,
                    'placeholders': []
                })
                continue
        if sections:
            sections[-1]['placeholders'].append(idx)

    # 5. Populate sections in reverse order
    for sec in reversed(sections):
        sub = sec['subsatgas']
        sec_unit = sec['unit_name']
        placeholders = sec['placeholders']
        if not placeholders:
            continue

        matched_items = []
        for (k_sub, k_unit), u_items in items_by_unit.items():
            if k_sub == sub:
                if k_unit == sec_unit or k_unit in sec_unit or sec_unit in k_unit:
                    matched_items = u_items
                    break

        template_row_tr = t0.rows[placeholders[0]]._tr

        if not matched_items:
            first_row = t0.rows[placeholders[0]]
            seen = set()
            u_cells = [c for c in first_row.cells if not (c._tc in seen or seen.add(c._tc))]
            if len(u_cells) >= 8:
                set_cell_text(u_cells[0], '')
                set_cell_text(u_cells[1], 'Gakkum')
                set_cell_text(u_cells[2], '-')
                set_cell_text(u_cells[3], '-')
                set_cell_text(u_cells[4], '-')
                set_cell_text(u_cells[5], '-')
                set_cell_text(u_cells[6], '-')
                set_cell_text(u_cells[7], '-')
            for ph_idx in placeholders[1:]:
                ph_row = t0.rows[ph_idx]
                seen_ph = set()
                ph_cells = [c for c in ph_row.cells if not (c._tc in seen_ph or seen_ph.add(c._tc))]
                for c in ph_cells:
                    set_cell_text(c, '')
            continue

        item_idx = 0
        last_row_obj = None

        for ph_idx in placeholders:
            if item_idx < len(matched_items):
                item = matched_items[item_idx]
                row = t0.rows[ph_idx]
                seen = set()
                u_cells = [c for c in row.cells if not (c._tc in seen or seen.add(c._tc))]
                if len(u_cells) >= 8:
                    set_cell_text(u_cells[0], f"{item_idx + 1}.")
                    set_cell_text(u_cells[1], item.get('satgas') or 'Gakkum')
                    set_cell_text(u_cells[2], item.get('activity_time') or '-')
                    set_cell_text(u_cells[3], item.get('personnel_count') or '-')
                    set_cell_text(u_cells[4], item.get('location') or '-')
                    set_cell_text(u_cells[5], item.get('activity_desc') or '-')
                    set_cell_text(u_cells[6], item.get('results') or '-')
                    set_cell_text(u_cells[7], item.get('notes') or '-')
                last_row_obj = row
                item_idx += 1
            else:
                row = t0.rows[ph_idx]
                seen = set()
                u_cells = [c for c in row.cells if not (c._tc in seen or seen.add(c._tc))]
                for c in u_cells:
                    set_cell_text(c, '')

        while item_idx < len(matched_items):
            item = matched_items[item_idx]
            new_tr = copy.deepcopy(template_row_tr)
            last_row_obj._tr.addnext(new_tr)
            
            new_row = None
            for r in t0.rows:
                if r._tr == new_tr:
                    new_row = r
                    break

            if new_row:
                seen = set()
                u_cells = [c for c in new_row.cells if not (c._tc in seen or seen.add(c._tc))]
                if len(u_cells) >= 8:
                    set_cell_text(u_cells[0], f"{item_idx + 1}.")
                    set_cell_text(u_cells[1], item.get('satgas') or 'Gakkum')
                    set_cell_text(u_cells[2], item.get('activity_time') or '-')
                    set_cell_text(u_cells[3], item.get('personnel_count') or '-')
                    set_cell_text(u_cells[4], item.get('location') or '-')
                    set_cell_text(u_cells[5], item.get('activity_desc') or '-')
                    set_cell_text(u_cells[6], item.get('results') or '-')
                    set_cell_text(u_cells[7], item.get('notes') or '-')
                last_row_obj = new_row
            item_idx += 1

    # 6. Table 1: Kejadian Menonjol
    if len(doc.tables) > 1:
        t1 = doc.tables[1]
        if incidents:
            for idx, inc in enumerate(incidents):
                if idx + 1 < len(t1.rows):
                    r = t1.rows[idx + 1]
                else:
                    new_tr = copy.deepcopy(t1.rows[1]._tr)
                    t1.rows[-1]._tr.addnext(new_tr)
                    r = t1.rows[-1]
                seen = set()
                u_cells = [c for c in r.cells if not (c._tc in seen or seen.add(c._tc))]
                if len(u_cells) >= 7:
                    set_cell_text(u_cells[0], str(idx + 1))
                    set_cell_text(u_cells[1], inc.get('incident_time') or '-')
                    set_cell_text(u_cells[2], inc.get('location') or '-')
                    set_cell_text(u_cells[3], inc.get('incident_desc') or '-')
                    set_cell_text(u_cells[4], inc.get('casualties_losses') or '-')
                    set_cell_text(u_cells[5], inc.get('actions_taken') or '-')
                    set_cell_text(u_cells[6], inc.get('notes') or '-')
        else:
            if len(t1.rows) > 1:
                r1 = t1.rows[1]
                seen = set()
                u_cells = [c for c in r1.cells if not (c._tc in seen or seen.add(c._tc))]
                if len(u_cells) >= 7:
                    set_cell_text(u_cells[0], '1')
                    set_cell_text(u_cells[1], 'NIHIL')
                    set_cell_text(u_cells[2], 'NIHIL')
                    set_cell_text(u_cells[3], 'NIHIL')
                    set_cell_text(u_cells[4], 'NIHIL')
                    set_cell_text(u_cells[5], 'NIHIL')
                    set_cell_text(u_cells[6], 'NIHIL')

    # 7. Post-processing: Ensure ALL table cells across all tables are Arial 12pt and rows cannot split
    from docx.oxml import OxmlElement
    for table in doc.tables:
        for row in table.rows:
            trPr = row._tr.get_or_add_trPr()
            if not row._tr.xpath('./w:trPr/w:cantSplit'):
                trPr.append(OxmlElement('w:cantSplit'))
            seen_cells = set()
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(12.0)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

