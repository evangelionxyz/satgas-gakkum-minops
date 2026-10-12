import os
import re
import copy
import docx
from docx.shared import Pt
from test_docx_engine import parse_docx, format_indonesian_date, OFFICIAL_UNITS

def clean_txt(t):
    return re.sub(r'\s+', ' ', t).strip().upper()

def set_cell_text(cell, text, bold=False, font_size=8.5, font_name='Arial'):
    p = cell.paragraphs[0]
    p.text = ''
    run = p.add_run(str(text or ''))
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold

def merge_items_to_master(template_path: str, output_path: str, items: list[dict], incidents: list[dict], target_date: str):
    doc = docx.Document(template_path)

    # 1. Update Title Date in Paragraphs
    formatted_date_upper = format_indonesian_date(target_date).upper()
    for p in doc.paragraphs[:20]:
        if "PERIODE HARI/TANGGAL" in p.text.upper():
            p.text = f"PERIODE HARI/TANGGAL : {formatted_date_upper}"
            # style font
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9.5)
                r.bold = True

    # 2. Update Date in Signature Block
    # e.g. "Jayapura, 01 Oktober 2026"
    for p in doc.paragraphs[-15:]:
        if re.search(r'JAYAPURA,\s+\d{1,2}\s+\w+\s+\d{4}', p.text.upper()):
            parts = formatted_date_upper.split(',')
            date_only = parts[-1].strip() if len(parts) > 1 else formatted_date_upper
            p.text = f"Jayapura, {date_only.title()}"
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9.5)

    # 3. Group items by (subsatgas, unit_name)
    items_by_unit = {}
    for it in items:
        sub = it.get('subsatgas', 'MINOPS').upper()
        unit = it.get('unit_name', '').upper().replace('–', '-').replace('—', '-')
        key = (sub, unit)
        items_by_unit.setdefault(key, []).append(it)

    # 4. Map template Table 0 sections
    t0 = doc.tables[0]
    sections = []
    current_sub = 'MINOPS'

    for idx, r in enumerate(t0.rows):
        seen = set()
        unique = [clean_txt(c.text) for c in r.cells if not (c._tc in seen or seen.add(c._tc))]
        non_empty = [c for c in unique if c]
        if len(non_empty) == 1:
            txt = non_empty[0]
            if txt in ['MINOPS', 'SUBSATGAS LIDIK', 'SUBSATGAS INVESTIGASI', 'SUBSATGAS ANALIS']:
                current_sub = txt.replace('SUBSATGAS ', '')
                continue
            elif any(k in txt for k in ['UNIT', 'WILAYAH']):
                sections.append({
                    'subsatgas': current_sub,
                    'unit_name': txt.replace('–', '-').replace('—', '-'),
                    'heading_idx': idx,
                    'placeholders': []
                })
                continue
        if sections:
            sections[-1]['placeholders'].append(idx)

    # 5. Populate sections in reverse order so inserted rows don't shift prior indices
    for sec in reversed(sections):
        sub = sec['subsatgas']
        sec_unit = sec['unit_name']
        placeholders = sec['placeholders']
        if not placeholders:
            continue

        # Find items for this section
        matched_items = []
        for (k_sub, k_unit), u_items in items_by_unit.items():
            if k_sub == sub:
                # Compare unit names flexible match
                if k_unit == sec_unit or k_unit in sec_unit or sec_unit in k_unit:
                    matched_items = u_items
                    break

        template_row_tr = t0.rows[placeholders[0]]._tr

        if not matched_items:
            # Clear placeholders or set NIHIL/-
            first_row = t0.rows[placeholders[0]]
            seen = set()
            u_cells = [c for c in first_row.cells if not (c._tc in seen or seen.add(c._tc))]
            if len(u_cells) >= 8:
                set_cell_text(u_cells[0], '')
                set_cell_text(u_cells[1], 'Gakkum')
                set_cell_text(u_cells[2], '-')
                set_cell_text(u_cells[3], '-')
                set_cell_text(u_cells[4], '-')
                set_cell_text(u_cells[5], '-')
                set_cell_text(u_cells[6], '-')
                set_cell_text(u_cells[7], '-')
            # Clear remaining placeholders if any
            for ph_idx in placeholders[1:]:
                ph_row = t0.rows[ph_idx]
                seen_ph = set()
                ph_cells = [c for c in ph_row.cells if not (c._tc in seen_ph or seen_ph.add(c._tc))]
                for c in ph_cells:
                    set_cell_text(c, '')
            continue

        # We have matched items!
        # First, populate existing placeholders
        item_idx = 0
        last_row_obj = None

        for ph_idx in placeholders:
            if item_idx < len(matched_items):
                item = matched_items[item_idx]
                row = t0.rows[ph_idx]
                seen = set()
                u_cells = [c for c in row.cells if not (c._tc in seen or seen.add(c._tc))]
                if len(u_cells) >= 8:
                    set_cell_text(u_cells[0], f"{item_idx + 1}.")
                    set_cell_text(u_cells[1], item.get('satgas') or 'Gakkum')
                    set_cell_text(u_cells[2], item.get('activity_time') or '-')
                    set_cell_text(u_cells[3], item.get('personnel_count') or '-')
                    set_cell_text(u_cells[4], item.get('location') or '-')
                    set_cell_text(u_cells[5], item.get('activity_desc') or '-')
                    set_cell_text(u_cells[6], item.get('results') or '-')
                    set_cell_text(u_cells[7], item.get('notes') or '-')
                last_row_obj = row
                item_idx += 1
            else:
                # Extra placeholder: clear it
                row = t0.rows[ph_idx]
                seen = set()
                u_cells = [c for c in row.cells if not (c._tc in seen or seen.add(c._tc))]
                for c in u_cells:
                    set_cell_text(c, '')

        # If more items than placeholders, clone and append after the last row
        while item_idx < len(matched_items):
            item = matched_items[item_idx]
            new_tr = copy.deepcopy(template_row_tr)
            last_row_obj._tr.addnext(new_tr)
            
            # Find the new row in table
            new_row = None
            for r in t0.rows:
                if r._tr == new_tr:
                    new_row = r
                    break

            if new_row:
                seen = set()
                u_cells = [c for c in new_row.cells if not (c._tc in seen or seen.add(c._tc))]
                if len(u_cells) >= 8:
                    set_cell_text(u_cells[0], f"{item_idx + 1}.")
                    set_cell_text(u_cells[1], item.get('satgas') or 'Gakkum')
                    set_cell_text(u_cells[2], item.get('activity_time') or '-')
                    set_cell_text(u_cells[3], item.get('personnel_count') or '-')
                    set_cell_text(u_cells[4], item.get('location') or '-')
                    set_cell_text(u_cells[5], item.get('activity_desc') or '-')
                    set_cell_text(u_cells[6], item.get('results') or '-')
                    set_cell_text(u_cells[7], item.get('notes') or '-')
                last_row_obj = new_row
            item_idx += 1

    # 6. Table 1: Kejadian Menonjol
    if len(doc.tables) > 1:
        t1 = doc.tables[1]
        if incidents:
            for idx, inc in enumerate(incidents):
                if idx + 1 < len(t1.rows):
                    r = t1.rows[idx + 1]
                else:
                    new_tr = copy.deepcopy(t1.rows[1]._tr)
                    t1.rows[-1]._tr.addnext(new_tr)
                    r = t1.rows[-1]
                seen = set()
                u_cells = [c for c in r.cells if not (c._tc in seen or seen.add(c._tc))]
                if len(u_cells) >= 7:
                    set_cell_text(u_cells[0], str(idx + 1))
                    set_cell_text(u_cells[1], inc.get('incident_time') or '-')
                    set_cell_text(u_cells[2], inc.get('location') or '-')
                    set_cell_text(u_cells[3], inc.get('incident_desc') or '-')
                    set_cell_text(u_cells[4], inc.get('casualties_losses') or '-')
                    set_cell_text(u_cells[5], inc.get('actions_taken') or '-')
                    set_cell_text(u_cells[6], inc.get('notes') or '-')
        else:
            # NIHIL
            if len(t1.rows) > 1:
                r1 = t1.rows[1]
                seen = set()
                u_cells = [c for c in r1.cells if not (c._tc in seen or seen.add(c._tc))]
                if len(u_cells) >= 7:
                    set_cell_text(u_cells[0], '1')
                    set_cell_text(u_cells[1], 'NIHIL')
                    set_cell_text(u_cells[2], 'NIHIL')
                    set_cell_text(u_cells[3], 'NIHIL')
                    set_cell_text(u_cells[4], 'NIHIL')
                    set_cell_text(u_cells[5], 'NIHIL')
                    set_cell_text(u_cells[6], 'NIHIL')

    doc.save(output_path)
    print(f"Successfully generated merged LAPHAR DOCX at: {output_path}")

if __name__ == "__main__":
    from test_docx_engine import parse_docx
    sample_folder = r'd:\Dev\satgas-gakkum-minops\old-backup\laporan\example'
    all_items = []
    all_incidents = []

    for fname in [
        'MINOPS JAYAPURA.docx',
        'MINOPS TIMIKA.docx',
        'MINOPS JAYAWIJAYA.docx',
        'LAPHAR_2026-10-05_SUBSATGAS LIDIK.docx',
        'LAPHAR SUBSATGAS INVEST 5 OKTOBER 2026.docx',
        'LAPORAN HARIAN SUBSATGAS ANALIS  5 OKTOBER 2026.docx'
    ]:
        res = parse_docx(os.path.join(sample_folder, fname))
        all_items.extend(res['items'])
        all_incidents.extend(res['incidents'])

    template = r'd:\Dev\satgas-gakkum-minops\backend\templates\TEMPLATE_LAPHAR_MASTER.docx'
    output = r'd:\Dev\satgas-gakkum-minops\backend\tests\LAPHAR_MERGED_TEST.docx'
    merge_items_to_master(template, output, all_items, all_incidents, "2026-10-05")
