import os
import re
import io
from datetime import datetime
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MONTHS_ID = {
    'JANUARI': '01', 'FEBRUARI': '02', 'MARET': '03', 'APRIL': '04',
    'MEI': '05', 'JUNI': '06', 'JULI': '07', 'AGUSTUS': '08',
    'SEPTEMBER': '09', 'OKTOBER': '10', 'NOVEMBER': '11', 'DESEMBER': '12'
}

DAYS_ID = {
    0: 'SENIN', 1: 'SELASA', 2: 'RABU', 3: 'KAMIS', 4: 'JUMAT', 5: 'SABTU', 6: 'MINGGU'
}

OFFICIAL_UNITS = {
    "MINOPS": [
        "UNIT JAYAPURA",
        "UNIT TIMIKA",
        "UNIT LANNY JAYA – KRING JAYAWIJAYA"
    ],
    "LIDIK": [
        "UNIT JAYAPURA",
        "UNIT LANNY JAYA",
        "UNIT YAHUKIMO",
        "UNIT PEGUNUNGAN BINTANG",
        "UNIT MIMIKA",
        "UNIT NDUGA",
        "UNIT INTAN JAYA",
        "UNIT NABIRE",
        "UNIT PUNCAK JAYA",
        "UNIT PUNCAK",
        "UNIT PUNCAK KRING SINAK"
    ],
    "INVESTIGASI": [
        "UNIT JAYAPURA",
        "UNIT TIMIKA",
        "UNIT LANNY JAYA",
        "UNIT NABIRE",
        "UNIT PUNCAK / ILAGA",
        "UNIT NDUGA",
        "UNIT INTAN JAYA",
        "UNIT YAHUKIMO",
        "UNIT PEGUNUNGAN BINTANG",
        "UNIT PUNCAK JAYA"
    ],
    "ANALIS": [
        "WILAYAH NDUGA",
        "WILAYAH PUNCAK",
        "WILAYAH INTAN JAYA",
        "WILAYAH PEGUNUNGAN BINTANG",
        "WILAYAH YAHUKIMO",
        "WILAYAH MIMIKA",
        "WILAYAH LANNY JAYA",
        "WILAYAH PANIAI",
        "WILAYAH NABIRE",
        "WILAYAH JAYAPURA",
        "WILAYAH PUNCAK JAYA"
    ]
}

def clean_text(text: str) -> str:
    if not text:
        return ''
    t = text.replace('\xa0', ' ').replace('–', '-').replace('—', '-')
    t = re.sub(r'\s+', ' ', t)
    return t.strip()

def extract_date_from_text(text: str) -> str | None:
    match = re.search(r'(\d{1,2})\s+(JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+(\d{4})', text.upper())
    if match:
        day = match.group(1).zfill(2)
        month = MONTHS_ID[match.group(2)]
        year = match.group(3)
        return f"{year}-{month}-{day}"
    match_iso = re.search(r'(\d{4})-(\d{2})-(\d{2})', text)
    if match_iso:
        return match_iso.group(0)
    return None

def format_indonesian_date(date_str: str) -> str:
    try:
        dt = datetime.strptime(date_str, "%Y-%m-%d")
        day_name = DAYS_ID[dt.weekday()]
        month_keys = list(MONTHS_ID.keys())
        month_name = month_keys[dt.month - 1]
        return f"{day_name}, {dt.day:02d} {month_name} {dt.year}"
    except Exception:
        return date_str

def get_row_cells(row):
    seen = set()
    cells = []
    for c in row.cells:
        if c._tc not in seen:
            seen.add(c._tc)
            cells.append(clean_text(c.text))
    return cells

def detect_subsatgas(filename: str, doc: docx.Document) -> str:
    upper_fname = filename.upper()
    if "MINOPS" in upper_fname:
        return "MINOPS"
    if "LIDIK" in upper_fname:
        return "LIDIK"
    if "INVEST" in upper_fname:
        return "INVESTIGASI"
    if "ANALIS" in upper_fname:
        return "ANALIS"

    for p in doc.paragraphs[:10]:
        t = p.text.upper()
        if "SUBSATGAS ANALIS" in t:
            return "ANALIS"
        if "SUBSATGAS LIDIK" in t:
            return "LIDIK"
        if "SUBSATGAS INVEST" in t:
            return "INVESTIGASI"
        if "MINOPS" in t:
            return "MINOPS"

    return "MINOPS"

def normalize_unit_name(unit_raw: str, subsatgas: str) -> str:
    clean = clean_text(unit_raw).upper()
    clean = re.sub(r'^(?:[IVXLCDM]+|\d+)[.)]?\s*', '', clean)
    clean = re.sub(r'^(?:UNIT|WILAYAH)\s+', '', clean)
    clean = clean.replace('–', '-').replace('—', '-').strip()

    if 'PEGUBIN' in clean or 'PEGUNUNG BINTANG' in clean or 'PEGUNUNGAN BINTANG' in clean:
        clean = 'PEGUNUNGAN BINTANG'
    elif 'WAMENA' in clean or 'KRING JAYAWIJAYA' in clean:
        clean = 'LANNY JAYA - KRING JAYAWIJAYA'
    elif 'SINAK' in clean:
        clean = 'PUNCAK KRING SINAK'
    elif 'ILAGA' in clean or 'PUNCAK / ILAGA' in clean:
        clean = 'PUNCAK / ILAGA'

    prefix = "WILAYAH " if subsatgas == "ANALIS" else "UNIT "
    formatted = prefix + clean

    # Match against official units list if available
    official_list = OFFICIAL_UNITS.get(subsatgas, [])
    for off in official_list:
        off_clean = off.replace('–', '-').replace('—', '-').upper()
        if formatted == off_clean:
            return off
        if clean in off_clean:
            return off

    return formatted

def parse_docx(file_path: str):
    doc = docx.Document(file_path)
    filename = os.path.basename(file_path)
    subsatgas = detect_subsatgas(filename, doc)
    
    detected_date = None
    # Check date in filename
    detected_date = extract_date_from_text(filename)
    if not detected_date:
        for p in doc.paragraphs[:15]:
            d = extract_date_from_text(p.text)
            if d:
                detected_date = d
                break

    items = []
    incidents = []

    if not doc.tables:
        return {"filename": filename, "subsatgas": subsatgas, "detected_date": detected_date, "items": items, "incidents": incidents}

    tbl = doc.tables[0]
    
    # Check if single unit file (e.g. MINOPS JAYAPURA)
    default_unit = None
    if subsatgas == "MINOPS":
        if "JAYAPURA" in filename.upper():
            default_unit = "UNIT JAYAPURA"
        elif "TIMIKA" in filename.upper():
            default_unit = "UNIT TIMIKA"
        elif "JAYAWIJAYA" in filename.upper() or "WAMENA" in filename.upper() or "LANNY" in filename.upper():
            default_unit = "UNIT LANNY JAYA – KRING JAYAWIJAYA"

    current_unit = default_unit
    last_context = {"time": "", "location": "", "personnel": ""}

    for r_idx, row in enumerate(tbl.rows):
        cells = get_row_cells(row)
        if not any(cells):
            continue

        # Check if unit header row
        # 1-cell or predominantly 1 cell filled
        non_empty = [c for c in cells if c]
        if len(non_empty) == 1:
            val = non_empty[0].upper()
            if any(k in val for k in ['UNIT', 'WILAYAH', 'KRING', 'JAYAPURA', 'TIMIKA', 'LANNY', 'YAHUKIMO', 'NDUGA', 'INTAN', 'NABIRE', 'PUNCAK', 'MIMIKA', 'PEGUNUNG']):
                # Header row
                current_unit = normalize_unit_name(non_empty[0], subsatgas)
                last_context = {"time": "", "location": "", "personnel": ""}
                continue

        # Check for noise rows
        row_str = ' '.join(cells).upper()
        if 'NO' in cells[0].upper() and ('WAKTU' in row_str or 'SATGAS' in row_str or 'JAM' in row_str or 'PUKUL' in row_str):
            continue
        if row_str.replace(' ', '') in ['1234567', '12345678']:
            continue
        if 'KEPOLISIAN NEGARA' in row_str or 'SATUAN TUGAS' in row_str or 'KASUBSATGAS' in row_str:
            continue

        if len(cells) < 6:
            continue

        # Map row by subsatgas
        item = {
            "subsatgas": subsatgas,
            "unit_name": current_unit or "UNIT JAYAPURA",
            "satgas": "Gakkum",
            "activity_time": "",
            "personnel_count": "",
            "location": "",
            "activity_desc": "",
            "results": "",
            "notes": ""
        }

        if subsatgas == "MINOPS":
            # 8 cols: NO, SATGAS, JAM, KUAT PERSONEL, LOKASI, KEGIATAN, HASIL, KET
            item["satgas"] = cells[1] if len(cells) > 1 and cells[1] else "Gakkum"
            item["activity_time"] = cells[2] if len(cells) > 2 else ""
            item["personnel_count"] = cells[3] if len(cells) > 3 else ""
            item["location"] = cells[4] if len(cells) > 4 else ""
            item["activity_desc"] = cells[5] if len(cells) > 5 else ""
            item["results"] = cells[6] if len(cells) > 6 else ""
            item["notes"] = cells[7] if len(cells) > 7 else ""

        elif subsatgas == "LIDIK":
            # 7 cols: NO, WAKTU, PERS, TEMPAT, KEGIATAN, HASIL, KET
            item["activity_time"] = cells[1] if len(cells) > 1 else ""
            item["personnel_count"] = cells[2] if len(cells) > 2 else ""
            item["location"] = cells[3] if len(cells) > 3 else ""
            item["activity_desc"] = cells[4] if len(cells) > 4 else ""
            item["results"] = cells[5] if len(cells) > 5 else ""
            item["notes"] = cells[6] if len(cells) > 6 else ""

        elif subsatgas == "INVESTIGASI":
            # 7 cols: NO, WAKTU, KUAT PERS, LOKASI, KEGIATAN, HASIL, KET
            item["activity_time"] = cells[1] if len(cells) > 1 else ""
            item["personnel_count"] = cells[2] if len(cells) > 2 else ""
            item["location"] = cells[3] if len(cells) > 3 else ""
            item["activity_desc"] = cells[4] if len(cells) > 4 else ""
            item["results"] = cells[5] if len(cells) > 5 else ""
            item["notes"] = cells[6] if len(cells) > 6 else ""

        elif subsatgas == "ANALIS":
            # 7 cols: NO, PUKUL, TEMPAT, KEGIATAN, HASIL, KUAT PERS, KET
            item["activity_time"] = cells[1] if len(cells) > 1 else ""
            item["location"] = cells[2] if len(cells) > 2 else ""
            item["activity_desc"] = cells[3] if len(cells) > 3 else ""
            item["results"] = cells[4] if len(cells) > 4 else ""
            item["personnel_count"] = cells[5] if len(cells) > 5 else ""
            item["notes"] = cells[6] if len(cells) > 6 else ""

            # Continuation handling for Analis
            if item["activity_time"]:
                last_context["time"] = item["activity_time"]
            else:
                item["activity_time"] = last_context["time"]

            if item["location"]:
                last_context["location"] = item["location"]
            else:
                item["location"] = last_context["location"]

            if item["personnel_count"]:
                last_context["personnel"] = item["personnel_count"]
            else:
                item["personnel_count"] = last_context["personnel"]

        # Only add if has meaningful activity or result
        if item["activity_desc"] or item["results"]:
            items.append(item)

    # Check Table 1 for Kejadian Menonjol
    if len(doc.tables) > 1:
        inc_tbl = doc.tables[1]
        for r in inc_tbl.rows[1:]:
            c = get_row_cells(r)
            if len(c) >= 6 and (c[3] or c[2]): # kejadian or lokasi
                if c[3].strip() not in ['-', '', '=']:
                    incidents.append({
                        "incident_time": c[1] if len(c) > 1 else "",
                        "location": c[2] if len(c) > 2 else "",
                        "incident_desc": c[3] if len(c) > 3 else "",
                        "casualties_losses": c[4] if len(c) > 4 else "",
                        "actions_taken": c[5] if len(c) > 5 else "",
                        "notes": c[6] if len(c) > 6 else ""
                    })

    return {
        "filename": filename,
        "subsatgas": subsatgas,
        "detected_date": detected_date,
        "items": items,
        "incidents": incidents
    }

print("Running test on all sample files...")
sample_folder = r'd:\Dev\satgas-gakkum-minops\old-backup\laporan\example'
all_parsed_items = []
all_parsed_incidents = []

for fname in [
    'MINOPS JAYAPURA.docx',
    'MINOPS TIMIKA.docx',
    'MINOPS JAYAWIJAYA.docx',
    'LAPHAR_2026-10-05_SUBSATGAS LIDIK.docx',
    'LAPHAR SUBSATGAS INVEST 5 OKTOBER 2026.docx',
    'LAPORAN HARIAN SUBSATGAS ANALIS  5 OKTOBER 2026.docx'
]:
    p = os.path.join(sample_folder, fname)
    res = parse_docx(p)
    print(f"File {fname}: Subsatgas={res['subsatgas']}, Date={res['detected_date']}, Items={len(res['items'])}, Incidents={len(res['incidents'])}")
    all_parsed_items.extend(res['items'])
    all_parsed_incidents.extend(res['incidents'])

print(f"Total parsed items across all files: {len(all_parsed_items)}")
