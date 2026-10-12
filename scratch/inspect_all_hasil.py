import glob
import os
import docx

hasil_files = glob.glob('d:/Dev/satgas-gakkum-minops/old-backup/laporan/hasil/*.docx')

for f in hasil_files:
    fname = os.path.basename(f)
    doc = docx.Document(f)
    t_drawings = 0
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if '<w:drawing' in cell._tc.xml:
                    t_drawings += 1
    has_dok = any('DOKUMENTASI' in p.text.upper() for p in doc.paragraphs)
    print(f"{fname}: t_drawings={t_drawings}, has_dok_section={has_dok}")
