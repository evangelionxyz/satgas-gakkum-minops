import glob
import os
import docx

hasil_files = glob.glob('d:/Dev/satgas-gakkum-minops/old-backup/laporan/hasil/*.docx')
print(f"Found {len(hasil_files)} hasil files")

for f in hasil_files[:5]:
    fname = os.path.basename(f)
    doc = docx.Document(f)
    p_drawings = sum(1 for p in doc.paragraphs if '<w:drawing' in p._p.xml)
    t_drawings = 0
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if '<w:drawing' in cell._tc.xml:
                    t_drawings += 1
    print(f"{fname}:")
    print(f"  Paragraphs: {len(doc.paragraphs)}, p_drawings: {p_drawings}")
    print(f"  Tables: {len(doc.tables)}, t_drawings: {t_drawings}")
    # If p_drawings > 0, check which paragraphs
    if p_drawings > 0:
        for idx, p in enumerate(doc.paragraphs):
            if '<w:drawing' in p._p.xml:
                print(f"    p[{idx}] text='{p.text[:60]}'")
