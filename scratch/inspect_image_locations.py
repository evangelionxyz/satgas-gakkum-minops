import docx

doc = docx.Document('d:/Dev/satgas-gakkum-minops/frontend/src-tauri/test/example/LAPHAR_2026-10-05_SUBSATGAS LIDIK.docx')
print(f"Paragraphs count: {len(doc.paragraphs)}")
print(f"Tables count: {len(doc.tables)}")

# Check paragraphs for drawings or text mentioning LAMPIRAN / DOKUMENTASI
for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    has_drawing = '<w:drawing' in xml
    text = p.text.strip()
    if has_drawing or 'LAMPIRAN' in text.upper() or 'DOKUMENTASI' in text.upper():
        print(f"Paragraph {i}: [has_drawing={has_drawing}] text='{text[:80]}'")

# Check tables for drawings
for t_idx, tbl in enumerate(doc.tables):
    drawings_in_table = 0
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            if '<w:drawing' in cell._tc.xml:
                drawings_in_table += 1
    print(f"Table {t_idx} ({len(tbl.rows)} rows x {len(tbl.columns)} cols): {drawings_in_table} cells with drawing")
