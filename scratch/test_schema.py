import zipfile
import re
import html
import docx

def xml_escape(s):
    return html.escape(str(s if s is not None else ''))

def set_tc_text(tc_xml, text, bold=False):
    tcPr_m = re.search(r'(<w:tcPr\b[^>]*>.*?</w:tcPr>)', tc_xml, flags=re.DOTALL)
    tcPr = tcPr_m.group(1) if tcPr_m else '<w:tcPr/>'

    pPr_m = re.search(r'(<w:pPr\b[^>]*>.*?</w:pPr>)', tc_xml, flags=re.DOTALL)
    pPr = pPr_m.group(1) if pPr_m else '<w:pPr><w:pStyle w:val="TableParagraph"/></w:pPr>'

    b_tag = '<w:b/>' if bold else ''
    esc_t = xml_escape(text)
    space_attr = ' xml:space="preserve"' if esc_t.startswith(' ') or esc_t.endswith(' ') else ''

    run = f'<w:r><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>{b_tag}<w:sz w:val="24"/><w:szCs w:val="24"/><w:lang w:val="id-ID"/></w:rPr><w:t{space_attr}>{esc_t}</w:t></w:r>'
    return f'<w:tc>{tcPr}<w:p>{pPr}{run}</w:p></w:tc>'

def set_row_values(tr_xml, values):
    tc_matches = list(re.finditer(r'<w:tc\b[^>]*>.*?</w:tc>', tr_xml, flags=re.DOTALL))
    if not tc_matches:
        return tr_xml

    new_tr = tr_xml[:tc_matches[0].start()]
    for i, m in enumerate(tc_matches):
        val = values[i] if i < len(values) else ""
        new_tr += set_tc_text(m.group(0), val)
    new_tr += tr_xml[tc_matches[-1].end():]
    return new_tr

def test():
    template_path = r'd:\Dev\satgas-gakkum-minops\backend\templates\TEMPLATE_LAPHAR_MASTER.docx'
    with zipfile.ZipFile(template_path) as z:
        entries = {info.filename: z.read(info.filename) for info in z.infolist()}

    doc_xml = entries['word/document.xml'].decode('utf-8')

    # Update title
    doc_xml = re.sub(
        r'(<w:t[^>]*>HARI\s*/\s*TANGGAL\s*:\s*).*?(</w:t>)',
        r'\g<1>SENIN, 12 OKTOBER 2026\2',
        doc_xml
    )

    # Locate Table 0
    t0_start = doc_xml.find('<w:tbl')
    t0_end = doc_xml.find('</w:tbl>', t0_start) + len('</w:tbl>')
    t0_xml = doc_xml[t0_start:t0_end]

    rows = re.findall(r'<w:tr\b[^>]*>.*?</w:tr>', t0_xml, flags=re.DOTALL)
    print(f"Table 0 has {len(rows)} rows.")

    # Test modifying row 3 (placeholder under Jayapura)
    new_r3 = set_row_values(rows[3], ["1.", "Gakkum", "08.00 WIT", "10 Pers", "Posko Jayapura", "Apel pagi kesiapan.", "Siap operasional.", "Aman."])
    rows[3] = new_r3

    # Reconstruct Table 0
    tblPr_match = re.search(r'^<w:tbl\b[^>]*>.*?(?=<w:tr)', t0_xml, flags=re.DOTALL)
    prefix = tblPr_match.group(0) if tblPr_match else '<w:tbl>'
    reconstructed_t0 = prefix + ''.join(rows) + '</w:tbl>'

    doc_xml = doc_xml[:t0_start] + reconstructed_t0 + doc_xml[t0_end:]
    entries['word/document.xml'] = doc_xml.encode('utf-8')

    out_docx = r'd:\Dev\satgas-gakkum-minops\scratch\test_schema_output.docx'
    import io
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)

    with open(out_docx, 'wb') as f:
        f.write(buf.getvalue())

    d = docx.Document(out_docx)
    print("Success loading output:", len(d.paragraphs), "Tables:", len(d.tables))
    t0 = d.tables[0]
    print("Row 3 cell 5 text:", t0.rows[3].cells[5].text)

if __name__ == '__main__':
    test()
