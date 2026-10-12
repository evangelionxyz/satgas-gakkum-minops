import docx

doc = docx.Document('d:/Dev/satgas-gakkum-minops/old-backup/laporan/hasil/LAPHAR MINOPS 05 OKTOBER 2026.docx')
for i in range(50, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    print(f"p[{i}]: text='{p.text}', xml_len={len(p._p.xml)}, has_drawing={'<w:drawing' in p._p.xml}")
