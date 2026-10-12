import glob
import os
import docx

files = glob.glob('d:/Dev/satgas-gakkum-minops/frontend/src-tauri/test/example/*.docx')

for f in sorted(files):
    fname = os.path.basename(f)
    doc = docx.Document(f)
    print(f"\n=== {fname} ===")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip().upper()
        if 'DOKUMENTASI' in t or 'LAMPIRAN' in t or 'FOTO' in t:
            print(f"  Paragraph {i}: '{p.text.strip()}' (has_drawing={'<w:drawing' in p._p.xml})")
    for t_idx, tbl in enumerate(doc.tables):
        for r_idx, row in enumerate(tbl.rows):
            for c_idx, cell in enumerate(row.cells):
                t = cell.text.strip().upper()
                if 'DOKUMENTASI' in t or 'LAMPIRAN' in t or 'FOTO' in t:
                    print(f"  Table {t_idx} Row {r_idx} Col {c_idx}: '{cell.text.strip()[:60]}' (has_drawing={'<w:drawing' in cell._tc.xml})")
