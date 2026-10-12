import zipfile
import re
import docx

master_path = 'd:/Dev/satgas-gakkum-minops/frontend/src-tauri/templates/TEMPLATE_LAPHAR_MASTER.docx'
minops_path = 'd:/Dev/satgas-gakkum-minops/frontend/src-tauri/test/example/MINOPS JAYAPURA.docx'

with zipfile.ZipFile(minops_path) as z_src:
    src_doc = z_src.read('word/document.xml').decode('utf-8')
    src_rels = z_src.read('word/_rels/document.xml.rels').decode('utf-8')
    src_media = {name: z_src.read(name) for name in z_src.namelist() if name.startswith('word/media/')}

re_rel = re.findall(r'<Relationship\b[^>]*Id="([^"]+)"[^>]*Type="[^"]*image"[^>]*Target="([^"]+)"', src_rels)
rel_map = {r_id: target for r_id, target in re_rel}

# Read master template
with zipfile.ZipFile(master_path) as z_m:
    master_entries = {name: z_m.read(name) for name in z_m.namelist()}

m_doc_xml = master_entries['word/document.xml'].decode('utf-8')
m_rels_xml = master_entries['word/_rels/document.xml.rels'].decode('utf-8')

# Get max rId in master rels
max_rid = max([int(num) for num in re.findall(r'rId(\d+)', m_rels_xml)] + [100])
next_rid = max_rid + 1

# Extract source table rows and drawings
src_tables = re.findall(r'<w:tbl\b[^>]*>.*?</w:tbl>', src_doc, re.DOTALL)
src_rows = re.findall(r'<w:tr\b[^>]*>.*?</w:tr>', src_tables[0], re.DOTALL)

new_rels = []
# Find row 3 drawing in src_rows
row3_drawings = re.findall(r'<w:drawing\b[^>]*>.*?</w:drawing>', src_rows[3], re.DOTALL)
print("Row 3 drawings count:", len(row3_drawings))

if row3_drawings:
    d = row3_drawings[0]
    # find old embed
    old_embed = re.search(r'r:embed="([^"]+)"', d).group(1)
    target = rel_map[old_embed]
    img_data = src_media['word/' + target]
    
    new_rid_str = f"rIdImg{next_rid}"
    new_media_name = f"merged_image_{next_rid}.jpeg"
    master_entries[f"word/media/{new_media_name}"] = img_data
    new_rels.append(f'<Relationship Id="{new_rid_str}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/{new_media_name}"/>')
    
    # replace embed in drawing
    updated_drawing = re.sub(r'r:embed="[^"]*"', f'r:embed="{new_rid_str}"', d)
    updated_drawing = re.sub(r'id="\d+"', f'id="{next_rid + 50000}"', updated_drawing, count=1)
    next_rid += 1
    
    # inject into master row 3 col 7 (KET)
    m_tables = re.findall(r'<w:tbl\b[^>]*>.*?</w:tbl>', m_doc_xml, re.DOTALL)
    m_rows = re.findall(r'<w:tr\b[^>]*>.*?</w:tr>', m_tables[0], re.DOTALL)
    r3 = m_rows[3]
    tcs = re.findall(r'<w:tc\b[^>]*>.*?</w:tc>', r3, re.DOTALL)
    # replace col 7
    old_tc7 = tcs[7]
    tc_pr = re.search(r'<w:tcPr\b[^>]*>.*?</w:tcPr>', old_tc7, re.DOTALL).group(0)
    new_tc7 = f'<w:tc>{tc_pr}<w:p><w:pPr><w:pStyle w:val="TableParagraph"/></w:pPr><w:r><w:rPr><w:noProof/></w:rPr>{updated_drawing}</w:r></w:p></w:tc>'
    new_r3 = r3.replace(old_tc7, new_tc7, 1)
    m_doc_xml = m_doc_xml.replace(r3, new_r3, 1)

master_entries['word/document.xml'] = m_doc_xml.encode('utf-8')
if new_rels:
    updated_rels_str = m_rels_xml.replace('</Relationships>', "".join(new_rels) + '</Relationships>')
    master_entries['word/_rels/document.xml.rels'] = updated_rels_str.encode('utf-8')

# Write output test
proto_out = 'd:/Dev/satgas-gakkum-minops/scratch/test_proto_out.docx'
with zipfile.ZipFile(proto_out, 'w', zipfile.ZIP_DEFLATED) as z_out:
    for name, data in master_entries.items():
        z_out.writestr(name, data)

# Verify with python-docx
doc = docx.Document(proto_out)
print(f"python-docx verified: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables")
cell = doc.tables[0].rows[3].cells[7]
print(f"Cell has drawing: {'<w:drawing' in cell._tc.xml}")
print("SUCCESS!")
