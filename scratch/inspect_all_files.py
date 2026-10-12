import os
import glob
import docx

example_dir = 'd:/Dev/satgas-gakkum-minops/frontend/src-tauri/test/example'
files = glob.glob(os.path.join(example_dir, '*.docx'))

for f in sorted(files):
    fname = os.path.basename(f)
    try:
        doc = docx.Document(f)
        p_drawings = sum(1 for p in doc.paragraphs if '<w:drawing' in p._p.xml)
        t_info = []
        for t_idx, tbl in enumerate(doc.tables):
            cells_with_drawing = 0
            for r_idx, row in enumerate(tbl.rows):
                for c_idx, cell in enumerate(row.cells):
                    if '<w:drawing' in cell._tc.xml:
                        cells_with_drawing += 1
            if cells_with_drawing > 0:
                t_info.append(f"Tbl{t_idx}: {cells_with_drawing} cells")
        print(f"{fname}:")
        print(f"  Paragraphs: {len(doc.paragraphs)} (drawings: {p_drawings})")
        print(f"  Tables: {len(doc.tables)}, drawing tables: {', '.join(t_info) if t_info else 'none'}")
    except Exception as e:
        print(f"{fname}: ERROR {e}")
